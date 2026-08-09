import re
import unicodedata
from pathlib import Path

from PIL import Image, ImageOps

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ============================================================
# Section names
# ============================================================

ZH_SECTION_NAMES = [
    (
        "executive_summary",
        "一、报告摘要",
    ),
    (
        "key_disaster_indicators",
        "二、核心灾情指标",
    ),
    (
        "regional_assessment",
        "三、分区评估结果",
    ),
    (
        "evidence_support_and_consistency_check",
        "四、证据支撑与一致性校验",
    ),
    (
        "limitations_and_nonconclusive_items",
        "五、证据局限与不可下结论事项",
    ),
]


EN_SECTION_NAMES = [
    (
        "executive_summary",
        "1. Executive Summary",
    ),
    (
        "key_disaster_indicators",
        "2. Key Disaster Indicators",
    ),
    (
        "regional_assessment",
        "3. Regional Assessment",
    ),
    (
        "evidence_support_and_consistency_check",
        "4. Evidence Support and Consistency Check",
    ),
    (
        "limitations_and_nonconclusive_items",
        "5. Limitations and Non-conclusive Items",
    ),
]


# ============================================================
# Text cleanup
# ============================================================

def sanitize_text(value):
    """
    Remove problematic Unicode noncharacters/control characters
    that may produce strange display results in Word/WPS.
    """

    text = str(value or "")

    text = unicodedata.normalize(
        "NFC",
        text,
    )

    # Remove Unicode noncharacters occasionally produced
    # by generated text / copy workflows.
    text = (
        text
        .replace("\ufffe", "")
        .replace("\uffff", "")
    )

    # Normalize several dash variants.
    text = (
        text
        .replace("\u2010", "-")
        .replace("\u2011", "-")
    )

    # Keep newline/tab but remove other C0 controls.
    text = re.sub(
        r"[\x00-\x08\x0b\x0c\x0e-\x1f]",
        "",
        text,
    )

    return text.strip()


def lang_text(obj, lang):
    if not isinstance(obj, dict):
        return sanitize_text(obj)

    return sanitize_text(
        obj.get(lang, "")
    )


# ============================================================
# Font control
# ============================================================

def set_run_font(
    run,
    lang,
    size=None,
    bold=None,
    color=None,
):
    """
    Explicitly write OOXML rFonts.
    This is critical for Chinese Word/WPS compatibility.
    """

    if lang == "zh-CN":
        font = "Microsoft YaHei"
    else:
        font = "Times New Roman"

    run.font.name = font

    rPr = (
        run._element
        .get_or_add_rPr()
    )

    rFonts = rPr.rFonts

    if rFonts is None:
        rFonts = OxmlElement(
            "w:rFonts"
        )
        rPr.insert(0, rFonts)

    rFonts.set(
        qn("w:ascii"),
        font,
    )

    rFonts.set(
        qn("w:hAnsi"),
        font,
    )

    rFonts.set(
        qn("w:eastAsia"),
        (
            "Microsoft YaHei"
            if lang == "zh-CN"
            else "Times New Roman"
        ),
    )

    if size is not None:
        run.font.size = Pt(size)

    if bold is not None:
        run.bold = bold

    if color is not None:
        run.font.color.rgb = RGBColor(
            *color
        )


def configure_styles(
    doc,
    lang,
):
    """
    Do not rely on Word theme fonts.
    Force explicit fonts for every main style.
    """

    font = (
        "Microsoft YaHei"
        if lang == "zh-CN"
        else "Times New Roman"
    )

    style_names = [
        "Normal",
        "Title",
        "Heading 1",
        "Heading 2",
        "List Bullet",
    ]

    for style_name in style_names:

        try:
            style = doc.styles[
                style_name
            ]
        except KeyError:
            continue

        style.font.name = font

        rPr = (
            style._element
            .get_or_add_rPr()
        )

        rFonts = rPr.rFonts

        if rFonts is None:
            rFonts = OxmlElement(
                "w:rFonts"
            )
            rPr.insert(
                0,
                rFonts,
            )

        rFonts.set(
            qn("w:ascii"),
            font,
        )

        rFonts.set(
            qn("w:hAnsi"),
            font,
        )

        rFonts.set(
            qn("w:eastAsia"),
            (
                "Microsoft YaHei"
                if lang == "zh-CN"
                else "Times New Roman"
            ),
        )


# ============================================================
# Image compatibility
# ============================================================

def normalize_image(
    source,
    compat_dir,
    asset_id,
):
    """
    Convert every visual asset into a standard RGB PNG.

    Avoid palette / alpha / unusual PNG-mode compatibility
    problems in Word or WPS.
    """

    source = Path(source)

    if not source.is_file():
        return None

    compat_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    target = (
        compat_dir
        / f"{asset_id}.png"
    )

    try:
        with Image.open(
            source
        ) as img:

            img = ImageOps.exif_transpose(
                img
            )

            if img.mode != "RGB":
                img = img.convert(
                    "RGB"
                )

            img.save(
                target,
                format="PNG",
                optimize=False,
            )

    except Exception as e:
        print(
            "[IMAGE NORMALIZE FAIL]",
            source,
            e,
        )
        return None

    return target


# ============================================================
# Word helpers
# ============================================================

def add_text_paragraph(
    doc,
    text,
    lang,
    style=None,
    size=10.5,
    bold=False,
    center=False,
    after=5,
):
    text = sanitize_text(
        text
    )

    p = doc.add_paragraph(
        style=style
    )

    if center:
        p.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

    p.paragraph_format.space_after = Pt(
        after
    )

    run = p.add_run(
        text
    )

    set_run_font(
        run,
        lang,
        size=size,
        bold=bold,
    )

    return p


def add_heading_manual(
    doc,
    text,
    lang,
):
    p = doc.add_paragraph()

    p.paragraph_format.space_before = Pt(
        9
    )

    p.paragraph_format.space_after = Pt(
        6
    )

    run = p.add_run(
        sanitize_text(text)
    )

    set_run_font(
        run,
        lang,
        size=15,
        bold=True,
        color=(31, 78, 121),
    )

    return p


def add_meta_line(
    doc,
    item,
    lang,
):
    claim = sanitize_text(
        item.get(
            "claim_id",
            ""
        )
    )

    evidence = [
        sanitize_text(x)
        for x in item.get(
            "evidence_ids",
            []
        )
    ]

    parts = []

    if lang == "zh-CN":

        if claim:
            parts.append(
                f"声明编号：{claim}"
            )

        if evidence:
            parts.append(
                "证据编号："
                + ", ".join(
                    evidence
                )
            )

    else:

        if claim:
            parts.append(
                f"Claim: {claim}"
            )

        if evidence:
            parts.append(
                "Evidence: "
                + ", ".join(
                    evidence
                )
            )

    if not parts:
        return

    p = doc.add_paragraph()

    p.paragraph_format.space_before = Pt(
        0
    )

    p.paragraph_format.space_after = Pt(
        6
    )

    run = p.add_run(
        " | ".join(parts)
    )

    set_run_font(
        run,
        lang,
        size=8.5,
        color=(100, 100, 100),
    )


def remove_table_borders(
    table,
):
    tblPr = table._tbl.tblPr

    borders = OxmlElement(
        "w:tblBorders"
    )

    for name in [
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV",
    ]:
        edge = OxmlElement(
            f"w:{name}"
        )

        edge.set(
            qn("w:val"),
            "nil",
        )

        borders.append(
            edge
        )

    tblPr.append(
        borders
    )


# ============================================================
# Visual evidence
# ============================================================

def caption_for(
    asset,
    lang,
):
    return lang_text(
        asset.get(
            "caption",
            {}
        ),
        lang,
    )


def evidence_caption(
    asset,
    lang,
):
    ids = [
        sanitize_text(x)
        for x in asset.get(
            "evidence_ids",
            []
        )
    ]

    if not ids:
        return ""

    if lang == "zh-CN":
        return (
            "证据编号："
            + ", ".join(ids)
        )

    return (
        "Evidence: "
        + ", ".join(ids)
    )


def add_full_visual(
    doc,
    asset,
    lang,
    compat_dir,
):
    asset_id = sanitize_text(
        asset.get(
            "asset_id",
            "image"
        )
    )

    image = normalize_image(
        asset.get(
            "path",
            ""
        ),
        compat_dir,
        asset_id,
    )

    if image is None:
        return

    p = doc.add_paragraph()

    p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = p.add_run()

    run.add_picture(
        str(image),
        width=Inches(5.7),
    )

    add_text_paragraph(
        doc,
        caption_for(
            asset,
            lang,
        ),
        lang,
        size=9,
        center=True,
        after=1,
    )

    evidence = evidence_caption(
        asset,
        lang,
    )

    if evidence:
        p2 = add_text_paragraph(
            doc,
            evidence,
            lang,
            size=7.5,
            center=True,
            after=6,
        )

        for r in p2.runs:
            r.font.color.rgb = RGBColor(
                110,
                110,
                110,
            )


def add_gallery_visual(
    cell,
    asset,
    lang,
    compat_dir,
):
    cell.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.TOP
    )

    asset_id = sanitize_text(
        asset.get(
            "asset_id",
            "image"
        )
    )

    image = normalize_image(
        asset.get(
            "path",
            ""
        ),
        compat_dir,
        asset_id,
    )

    if image is None:
        return

    p = cell.paragraphs[0]

    p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    r = p.add_run()

    r.add_picture(
        str(image),
        width=Inches(2.55),
    )

    cap = cell.add_paragraph()

    cap.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    cr = cap.add_run(
        caption_for(
            asset,
            lang,
        )
    )

    set_run_font(
        cr,
        lang,
        size=8.5,
    )

    evidence = evidence_caption(
        asset,
        lang,
    )

    if evidence:

        ep = cell.add_paragraph()

        ep.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        er = ep.add_run(
            evidence
        )

        set_run_font(
            er,
            lang,
            size=7,
            color=(110, 110, 110),
        )


def add_visual_section(
    doc,
    visuals,
    lang,
    compat_dir,
):
    if not visuals:
        return

    title = (
        "图像证据"
        if lang == "zh-CN"
        else "Visual Evidence"
    )

    add_heading_manual(
        doc,
        title,
        lang,
    )

    full_width = []
    gallery = []

    for asset in visuals:

        if (
            asset.get(
                "asset_type"
            )
            == "pre_post_pair"
        ):
            full_width.append(
                asset
            )

        else:
            gallery.append(
                asset
            )

    for asset in full_width:

        add_full_visual(
            doc,
            asset,
            lang,
            compat_dir,
        )

    if not gallery:
        return

    rows = (
        len(gallery) + 1
    ) // 2

    table = doc.add_table(
        rows=rows,
        cols=2,
    )

    table.autofit = False

    remove_table_borders(
        table
    )

    for i, asset in enumerate(
        gallery
    ):

        row = i // 2
        col = i % 2

        add_gallery_visual(
            table.cell(
                row,
                col,
            ),
            asset,
            lang,
            compat_dir,
        )


# ============================================================
# Main Word renderer
# ============================================================

def render_word_language(
    report,
    visuals,
    output,
    lang,
):
    doc = Document()

    configure_styles(
        doc,
        lang,
    )

    section = doc.sections[0]

    section.top_margin = Inches(
        0.72
    )

    section.bottom_margin = Inches(
        0.72
    )

    section.left_margin = Inches(
        0.82
    )

    section.right_margin = Inches(
        0.82
    )

    if lang == "zh-CN":
        title = "遥感灾情初步研判报告"
        sections = ZH_SECTION_NAMES
    else:
        title = (
            "Preliminary Remote-Sensing "
            "Assessment Report"
        )
        sections = EN_SECTION_NAMES

    title_p = doc.add_paragraph()

    title_p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title_p.paragraph_format.space_after = Pt(
        10
    )

    title_run = title_p.add_run(
        title
    )

    set_run_font(
        title_run,
        lang,
        size=20,
        bold=True,
    )

    scene = sanitize_text(
        report.get(
            "task_info",
            {}
        ).get(
            "scene_uid",
            ""
        )
    )

    scene_label = (
        f"场景编号：{scene}"
        if lang == "zh-CN"
        else f"Scene UID: {scene}"
    )

    add_text_paragraph(
        doc,
        scene_label,
        lang,
        size=10,
        center=True,
        after=12,
    )

    report_sections = report.get(
        "sections",
        {}
    )

    for key, section_name in sections:

        add_heading_manual(
            doc,
            section_name,
            lang,
        )

        value = report_sections.get(
            key
        )

        # -----------------------------------------
        # List sections
        # -----------------------------------------

        if isinstance(
            value,
            list
        ):

            if not value:

                add_text_paragraph(
                    doc,
                    (
                        "无"
                        if lang == "zh-CN"
                        else "None"
                    ),
                    lang,
                )

                continue

            for item in value:

                # Findings / limitations
                if "text" in item:

                    text = lang_text(
                        item.get(
                            "text",
                            {}
                        ),
                        lang,
                    )

                    p = doc.add_paragraph(
                        style="List Bullet"
                    )

                    p.paragraph_format.space_after = Pt(
                        2
                    )

                    run = p.add_run(
                        text
                    )

                    set_run_font(
                        run,
                        lang,
                        size=10.5,
                    )

                    add_meta_line(
                        doc,
                        item,
                        lang,
                    )

                # Evidence consistency records
                else:

                    eid = sanitize_text(
                        item.get(
                            "evidence_id",
                            ""
                        )
                    )

                    claim_ids = [
                        sanitize_text(x)
                        for x in item.get(
                            "claim_ids",
                            []
                        )
                    ]

                    if lang == "zh-CN":
                        line = (
                            f"证据 {eid} → "
                            + ", ".join(
                                claim_ids
                            )
                        )
                    else:
                        line = (
                            f"Evidence {eid} → "
                            + ", ".join(
                                claim_ids
                            )
                        )

                    p = doc.add_paragraph(
                        style="List Bullet"
                    )

                    run = p.add_run(
                        line
                    )

                    set_run_font(
                        run,
                        lang,
                        size=9.5,
                    )

        # -----------------------------------------
        # Summary
        # -----------------------------------------

        else:

            add_text_paragraph(
                doc,
                lang_text(
                    value,
                    lang,
                ),
                lang,
                size=10.5,
                after=8,
            )

    # -----------------------------------------
    # Visual evidence
    # -----------------------------------------

    output = Path(
        output
    )

    compat_dir = (
        output.parent
        / "images_compat"
    )

    add_visual_section(
        doc,
        visuals,
        lang,
        compat_dir,
    )

    # -----------------------------------------
    # Disclaimer
    # -----------------------------------------

    disclaimer_title = (
        "声明"
        if lang == "zh-CN"
        else "Disclaimer"
    )

    add_heading_manual(
        doc,
        disclaimer_title,
        lang,
    )

    add_text_paragraph(
        doc,
        lang_text(
            report.get(
                "disclaimer",
                {}
            ),
            lang,
        ),
        lang,
        size=9.5,
        after=6,
    )

    doc.save(
        output
    )

    return output


# ============================================================
# Markdown
# ============================================================

def render_markdown_language(
    report,
    visuals,
    output,
    lang,
):
    if lang == "zh-CN":

        title = (
            "# 遥感灾情初步研判报告"
        )

        sections = (
            ZH_SECTION_NAMES
        )

        claim_label = "声明编号"

        evidence_label = "证据编号"

    else:

        title = (
            "# Preliminary Remote-Sensing "
            "Assessment Report"
        )

        sections = (
            EN_SECTION_NAMES
        )

        claim_label = "Claim"

        evidence_label = "Evidence"

    lines = [
        title,
        "",
    ]

    scene = sanitize_text(
        report.get(
            "task_info",
            {}
        ).get(
            "scene_uid",
            ""
        )
    )

    lines.append(
        (
            f"**场景编号：** {scene}"
            if lang == "zh-CN"
            else f"**Scene UID:** {scene}"
        )
    )

    report_sections = report.get(
        "sections",
        {}
    )

    for key, name in sections:

        lines.extend([
            "",
            f"## {name}",
            "",
        ])

        value = report_sections.get(
            key
        )

        if isinstance(
            value,
            list
        ):

            for item in value:

                if "text" in item:

                    lines.append(
                        "- "
                        + lang_text(
                            item.get(
                                "text",
                                {}
                            ),
                            lang,
                        )
                    )

                    meta = []

                    if item.get(
                        "claim_id"
                    ):

                        meta.append(
                            f"{claim_label}: "
                            + sanitize_text(
                                item[
                                    "claim_id"
                                ]
                            )
                        )

                    if item.get(
                        "evidence_ids"
                    ):

                        meta.append(
                            f"{evidence_label}: "
                            + ", ".join(
                                sanitize_text(x)
                                for x in item[
                                    "evidence_ids"
                                ]
                            )
                        )

                    if meta:
                        lines.append(
                            "  - "
                            + " | ".join(meta)
                        )

                else:

                    eid = sanitize_text(
                        item.get(
                            "evidence_id",
                            ""
                        )
                    )

                    claims = ", ".join(
                        sanitize_text(x)
                        for x in item.get(
                            "claim_ids",
                            []
                        )
                    )

                    lines.append(
                        f"- {evidence_label} {eid} → {claims}"
                    )

        else:

            lines.append(
                lang_text(
                    value,
                    lang,
                )
            )

    if visuals:

        lines.extend([
            "",
            (
                "## 图像证据"
                if lang == "zh-CN"
                else "## Visual Evidence"
            ),
            "",
        ])

        for asset in visuals:

            caption = caption_for(
                asset,
                lang,
            )

            lines.append(
                f"### {caption}"
            )

            lines.append(
                f"![{caption}]({asset.get('path', '')})"
            )

            evidence = evidence_caption(
                asset,
                lang,
            )

            if evidence:
                lines.append(
                    evidence
                )

            lines.append("")

    lines.extend([
        "",
        (
            "## 声明"
            if lang == "zh-CN"
            else "## Disclaimer"
        ),
        "",
        lang_text(
            report.get(
                "disclaimer",
                {}
            ),
            lang,
        ),
    ])

    Path(
        output
    ).write_text(
        "\n".join(lines),
        encoding="utf-8",
    )
